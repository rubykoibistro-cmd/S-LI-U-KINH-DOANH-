"""Kiểm thử các phép tính kế toán trên bộ dữ liệu mẫu (du_lieu_mau) – số kỳ vọng tính tay."""
from __future__ import annotations

import csv
import shutil
from datetime import date
from pathlib import Path

import pytest

from ruby_koi_ketoan import nhap_file
from ruby_koi_ketoan.bao_cao import lap_bao_cao, tinh_toan
from ruby_koi_ketoan.cong_no import kiem_soat_hoa_don, tinh_cong_no_ncc
from ruby_koi_ketoan.doanh_thu import doanh_thu_ngay
from ruby_koi_ketoan.du_lieu import nap_du_lieu, tao_mau
from ruby_koi_ketoan.lai_lo import lai_lo
from ruby_koi_ketoan.nhap_xuat_ton import tinh_nhap_xuat_ton, tong_hop_theo_nvl
from ruby_koi_ketoan.tien_ich import doc_ngay, doc_so

MAU = Path(__file__).resolve().parent.parent / "du_lieu_mau"
D1, D2, D3 = date(2026, 10, 1), date(2026, 10, 2), date(2026, 10, 3)


@pytest.fixture()
def dl():
    return nap_du_lieu(MAU)


@pytest.mark.parametrize("vao, ra", [
    ("11,825,000", 11825000), ("1.500.000", 1500000), ("1.500", 1500), ("0,5", 0.5), ("16,66", 16.66),
    (" - ", 0), ("(1.000)", -1000), ("1.234.567 đ", 1234567), (12.5, 12.5), ("", 0), ("0.500", 0.5),
])
def test_doc_so(vao, ra):
    assert doc_so(vao) == pytest.approx(ra)


def test_doc_so_luong_dau_cham_la_thap_phan():
    assert doc_so("2.002", so_luong=True) == pytest.approx(2.002)
    assert doc_so("2.002") == 2002


@pytest.mark.parametrize("vao, ra", [("29/09/2026", date(2026, 9, 29)), ("1.7.2026", date(2026, 7, 1)),
                                     ("2026-08-31", date(2026, 8, 31)), (46294, date(2026, 9, 29))])
def test_doc_ngay(vao, ra):
    assert doc_ngay(vao) == ra


def test_nhap_xuat_ton_va_hao_hut(dl):
    kho = tinh_nhap_xuat_ton(dl, D1, D3)
    t = tong_hop_theo_nvl(kho.dong)
    bo = t["BO-JBS"]
    # Bò: 5 kg @378k + nhập 10 kg @390k -> giá BQ 386k; bán 30 món × 0,15 = 4,5 kg; kiểm 10 kg
    assert bo["gia_bq"] == pytest.approx(386000)
    assert bo["xuat_ban"] == pytest.approx(4.5)
    assert bo["that_thoat"] == pytest.approx(0.5)
    assert bo["gia_tri_that_thoat"] == pytest.approx(193000)
    assert bo["ty_le_hao_hut"] == pytest.approx(193000 / (4.5 * 386000 + 193000))
    ca = t["CA-MU"]
    assert ca["huy"] == pytest.approx(0.2) and ca["that_thoat"] == pytest.approx(0.3)
    assert ca["ton_cuoi"] == pytest.approx(4)
    assert t["TOM-20"]["that_thoat"] == pytest.approx(0.1)


def test_gia_von_bang_mua_cong_giam_ton(dl):
    kho = tinh_nhap_xuat_ton(dl, D1, D3)
    pl = lai_lo(dl, kho, D1, D3)
    # Mua: sổ quỹ 880k + 300k, mua chịu 3,9tr. Giảm tồn = (1.890.000+654.000+420.000) − (3.860.000+876.571,43+189.000)
    giam_ton = (1890000 + 654000 + 420000) - (3860000 + 4 * (654000 + 880000) / 7 + 189000)
    assert pl["gia_von"]["tong"] == pytest.approx(880000 + 300000 + 3900000 + giam_ton)
    # Đối chiếu chiều ngược lại: tổng lượng dùng thực tế × giá + rau không theo dõi tồn
    dung = 5 * 386000 + 3.0 * (654000 + 880000) / 7 + 1.1 * 210000 + 300000
    assert pl["gia_von"]["tong"] == pytest.approx(dung)
    assert pl["doanh_thu_thuan"] == pytest.approx(13700000 - 210000)


def test_cong_no_ncc(dl):
    cn = tinh_cong_no_ncc(dl, D3)
    assert cn.tong_con_no() == pytest.approx(2900000)
    qua_han = tinh_cong_no_ncc(dl, date(2026, 10, 20))
    assert any("quá hạn 12 ngày" in c["noi_dung"] for c in qua_han.canh_bao)


def test_soat_hoa_don_bat_sai_thanh_tien(dl):
    cb = kiem_soat_hoa_don(dl, D1, D3)
    assert any("Chanh" in c["noi_dung"] and "Sai thành tiền" in c["noi_dung"] for c in cb)


def test_doi_soat_quy_tien_mat(dl):
    d = doanh_thu_ngay(dl, D2)
    assert d["tien_mat_tinh"] == pytest.approx(2000000)
    assert d["chenh_lech_quy"] == pytest.approx(200000)  # đếm 1,9tr − (2tr − chi rau 300k)
    assert doanh_thu_ngay(dl, D3)["chenh_lech_quy"] == pytest.approx(0)


def test_nvl_chua_co_ton_dau_khong_theo_doi(tmp_path):
    shutil.copytree(MAU, tmp_path / "d")
    (tmp_path / "d" / "ton_dau_ky.csv").write_text(
        "ngay,ma_nvl,so_luong,don_gia\n2026-10-01,BO-JBS,5,378000\n", encoding="utf-8")
    dl = nap_du_lieu(tmp_path / "d")
    kho = tinh_nhap_xuat_ton(dl, D1, D3)
    assert {d.ma_nvl for d in kho.dong} == {"BO-JBS"}
    assert any("CHƯA có số kiểm kê đầu kỳ" in c["noi_dung"] for c in kho.canh_bao)


def test_nvl_khong_dinh_luong_tinh_dung_theo_kiem_ke(tmp_path):
    shutil.copytree(MAU, tmp_path / "d")
    (tmp_path / "d" / "dinh_luong.csv").write_text("ma_mon,ten_mon,ma_nvl,luong\n", encoding="utf-8")
    dl = nap_du_lieu(tmp_path / "d")
    t = tong_hop_theo_nvl(tinh_nhap_xuat_ton(dl, D1, D3).dong)
    assert t["BO-JBS"]["that_thoat"] == 0
    assert t["BO-JBS"]["xuat_ban"] == pytest.approx(5.0)  # 5 + 10 − 10
    assert t["BO-JBS"]["ty_le_hao_hut"] is None


def test_thieu_so_quy_thi_canh_bao_loi_nhuan(tmp_path):
    shutil.copytree(MAU, tmp_path / "d")
    (tmp_path / "d" / "chi_phi.csv").write_text(
        "ngay,nhom_chi_phi,noi_dung,so_tien,hinh_thuc_tt,ky_phan_bo,ghi_chu\n", encoding="utf-8")
    kq = tinh_toan(nap_du_lieu(tmp_path / "d"), D3, D3)
    assert len(kq["do_day_du"]["thieu_chi_phi"]) == 3
    assert any(c["loai"] == "THIẾU SỐ LIỆU" for c in kq["canh_bao"])


def test_lap_bao_cao_ra_file(dl, tmp_path):
    kq = lap_bao_cao(dl, D3, D3, "ngay", tmp_path)
    assert kq["tep_md"].exists() and kq["tep_xlsx"].exists()
    assert "BÁO CÁO KẾ TOÁN NỘI BỘ NGÀY 03/10/2026" in kq["markdown"]


# ------------------------------------------------------------------ nhập file gốc
def _dsr_gia(tep: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "QTBH"
    ws.append(["STT", "NGÀY", "DOANH \nTHU", "CÁC KHOẢN GIẢM TRỪ", None, None, None, None, None, None, None,
               "THUẾ \nGTGT", "TỔNG PHẢI \nTHU", None, None, None, None, "TIỀN BÁN HÀNG CẦN PHẢI NỘP", "GHI\nCHÚ"])
    ws.append([None, None, None, "GIẢM GIÁ NHÂN VIÊN 20%/30%", "ANH QUANG", "CHI CHUNG", "CN ANH HÀ ",
               "CN KHÁCH LẺ/CN TỈNH ĐOÀN", "ANH QUANG TIẾP KHÁCH ", "TRỪ LUI", "CÔNG NỢ NHÂN VIÊN ", None, None,
               "CA 1", "CA 2", "CK CHỊ Quyên", "TỔNG CỘNG"])
    ws.append(["TỔNG", None, 1])
    ws.append([1, date(2026, 9, 1), 11825000, 156000, None, 70000, 35000, None, None, None, None, 1182500,
               11564000, 5550000, 530000, None, 6080000, 5484000])
    ws.append([2, date(2026, 9, 2), 3346000, None, None, None, None, None, None, None, None, 334600, 3346000,
               1536000, None, None, 850000, 2496000])
    ws.append([3, date(2026, 9, 3), 0, None, None, None, None, None, None, None, None, 0, 0, None, None, None,
               "#REF!", "#REF!"])
    wb.create_sheet("sổ chi")
    wb.save(tep)


def test_nhap_dsr_qtbh(tmp_path):
    tep = tmp_path / "dsr.xlsx"
    _dsr_gia(tep)
    tao_mau(tmp_path / "d")
    bao = nhap_file.nhap_dsr(tep, tmp_path / "d")
    assert any("02/09" in b and "TỔNG CỘNG" in b for b in bao)
    dl = nap_du_lieu(tmp_path / "d")
    d = doanh_thu_ngay(dl, date(2026, 9, 1))
    assert d["doanh_thu_thuan"] == 11825000 - 156000 - 70000
    assert d["cong_no_khach"] == 35000
    assert d["tien_mat_pos"] == d["tien_mat_tinh"] == 5484000
    assert len(dl["doanh_thu_ngay.csv"]) == 2  # ngày 03/09 doanh thu 0 bị bỏ
    nhap_file.nhap_dsr(tep, tmp_path / "d")  # chạy lại không nhân đôi
    assert len(nap_du_lieu(tmp_path / "d")["doanh_thu_ngay.csv"]) == 2


def _so_quy_gia(tep: Path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for ngay, dong in {1: [("PC02", "cô Hạnh", "K", 165000), ("PC03", "cám cá", "L", 500000),
                           ("PC04", "lương bảo vệ T7", "H", 8000000)],
                       2: [("PC02", "Thuận Phát", "AC", 5000000)]}.items():
        ws = wb.create_sheet(str(ngay))
        for _ in range(10):
            ws.append([])
        ws.append([None, None, "PT01", None, "DOANH THU", 1000000])
        for ct, dg, cot, tien in dong:
            r = ws.max_row + 1
            ws.cell(r, 4, ct)
            ws.cell(r, 5, dg)
            ws[f"{cot}{r}"] = tien
            if cot == "AC":
                ws[f"AG{r}"] = "dtt 9/9"
        ws.cell(ws.max_row + 1, 1, "CỘNG")
    ws = wb.create_sheet("TỔNG")
    ws["AG43"] = "=1000000+200000"
    wb.save(tep)


def test_nhap_so_quy(tmp_path):
    tep = tmp_path / "sq.xlsx"
    _so_quy_gia(tep)
    tao_mau(tmp_path / "d")
    bao = nhap_file.nhap_so_quy(tep, tmp_path / "d", date(2026, 8, 1))
    assert any("AG43" in b and "1,200,000" in b for b in bao)
    with (tmp_path / "d" / "chi_phi.csv").open(encoding="utf-8-sig") as f:
        cp = {r["noi_dung"]: r for r in csv.DictReader(f)}
    assert cp["cám cá"]["nhom_chi_phi"] == "van_hanh_ho_koi"
    assert cp["lương bảo vệ T7"]["ky_phan_bo"] == "2026-07"
    assert cp["cô Hạnh"]["nhom_chi_phi"] == "nvl_bep"
    with (tmp_path / "d" / "nhap_hang.csv").open(encoding="utf-8-sig") as f:
        no = list(csv.DictReader(f))
    assert len(no) == 1 and no[0]["hinh_thuc_tt"] == "cong_no" and no[0]["han_thanh_toan"] == "2026-09-09"
    # Lương T7 trả trong T8 không vào lãi lỗ T8
    dl = nap_du_lieu(tmp_path / "d")
    pl = lai_lo(dl, tinh_nhap_xuat_ton(dl, date(2026, 8, 1), date(2026, 8, 31)), date(2026, 8, 1), date(2026, 8, 31))
    assert "luong" not in pl["nhan_su"]
    assert pl["gia_von"]["tong"] == pytest.approx(165000 + 5000000)


def test_bang_dieu_hanh(dl):
    from ruby_koi_ketoan.bang_dieu_hanh import du_lieu_bang
    b = du_lieu_bang(dl, D3)
    tq = b["tong_quan"][0]
    assert tq["doanh_thu"] == 13490000 and tq["so_ngay"] == 3
    assert tq["cong_no_ncc"] == 2900000
    assert [r["ngay"] for r in b["ngay"]] == ["2026-10-01", "2026-10-02", "2026-10-03"]
    assert b["thang"][0]["du_so_quy"] is True and b["thang"][0]["ebitda"] is not None
    assert 1 <= len(b["viec"]) <= 5
