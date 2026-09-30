"""Nhập số liệu từ các file gốc đang dùng ở Ruby Koi sang file dữ liệu chuẩn.

Hỗ trợ:
  - DSR tháng (RUBYKOI.DSR.Tx.2026.xlsx): sheet QTBH (doanh thu, giảm trừ, thanh toán),
    sheet SALES (số khách, số bill), sheet PM (số lượng món bán theo ngày), sheet "sổ chi".
  - Sổ Quỹ tháng (SOQUYTx.2026.xlsx): 31 sheet ngày, mỗi dòng một phiếu chi theo cột nhóm.
  - Sổ mua hàng chi tiết (Sổ_Mua_Hàng_Chi_Phí_Hàng_Ngày).
  - Báo cáo bán hàng theo mặt hàng xuất từ iPOS (CSV/XLSX, 1 ngày).

Mọi lần nhập đều "thay thế theo ngày + nguồn": chạy lại cùng file không bị nhân đôi số.
"""
from __future__ import annotations

import csv
import re
from datetime import date
from pathlib import Path

from .du_lieu import BANG
from .tien_ich import chuan_hoa_ten, doc_ngay, doc_so, so_ngay_trong_thang


def _cot(ten_bang: str) -> list[str]:
    return [c.rstrip("*") for c in BANG[ten_bang][0]]


def loc_mua_chiu_trung(thu_muc: Path, dong_moi: list[dict], nguon: str) -> tuple[list[dict], list[str]]:
    """Một khoản mua chịu hay được ghi ở CẢ sổ chi lẫn cột CÔNG NỢ sổ quỹ. Bỏ dòng mới nếu nguồn
    khác đã có khoản nợ cùng ngày, cùng số tiền (±1.000đ)."""
    tep = Path(thu_muc) / "nhap_hang.csv"
    if not tep.exists():
        return dong_moi, []
    with tep.open(encoding="utf-8-sig", newline="") as f:
        da_co = [(doc_ngay(r["ngay"]), doc_so(r["thanh_tien"])) for r in csv.DictReader(f)
                 if r.get("hinh_thuc_tt") == "cong_no" and not (r.get("ghi_chu") or "").startswith(f"[{nguon}]")]
    giu, bo = [], []
    for r in dong_moi:
        if r.get("hinh_thuc_tt") == "cong_no" and any(
                n == doc_ngay(r["ngay"]) and abs(t - r["thanh_tien"]) <= 1000 for n, t in da_co):
            bo.append(f"  ↷ Bỏ khoản mua chịu {r['ncc']} {r['ngay']} {r['thanh_tien']:,.0f}đ – đã có ở nguồn khác (tránh ghi nợ 2 lần)")
        else:
            giu.append(r)
    return giu, bo


def ghi_bang(thu_muc: Path, ten_bang: str, dong_moi: list[dict], nguon: str) -> tuple[int, int]:
    """Ghi dòng mới vào bảng, xoá các dòng cũ cùng nguồn trong những ngày được nhập lại."""
    tep = Path(thu_muc) / ten_bang
    cot = _cot(ten_bang)
    cu = []
    if tep.exists():
        with tep.open(encoding="utf-8-sig", newline="") as f:
            cu = list(csv.DictReader(f))
    ngay_moi = {str(r["ngay"]) for r in dong_moi}
    the = f"[{nguon}]"

    def la_cua_nguon(r):
        if ten_bang in ("doanh_thu_ngay.csv", "ban_hang_pos.csv"):
            return True  # mỗi ngày chỉ có một nguồn POS
        return (r.get("ghi_chu") or "").startswith(the)

    giu = [r for r in cu if not (str(doc_ngay(r.get("ngay")) or r.get("ngay")) in ngay_moi and la_cua_nguon(r))]
    for r in dong_moi:
        if "ghi_chu" in cot and ten_bang not in ("doanh_thu_ngay.csv",):
            r["ghi_chu"] = f"{the} {r.get('ghi_chu', '')}".strip()
    tat_ca = giu + [{k: ("" if r.get(k) is None else r.get(k)) for k in cot} for r in dong_moi]
    tat_ca.sort(key=lambda r: str(doc_ngay(r.get("ngay")) or r.get("ngay")))
    with tep.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cot, extrasaction="ignore")
        w.writeheader()
        w.writerows(tat_ca)
    return len(dong_moi), len(cu) - len(giu)


def _so(v, so_luong: bool = False) -> float | None:
    if v is None or (isinstance(v, str) and (v.strip() == "" or v.strip().startswith("#"))):
        return None
    return doc_so(v, so_luong=so_luong)


def _tim_hang_tieu_de(ws, tu_khoa: list[str], toi_da: int = 30) -> int | None:
    for i, hang in enumerate(ws.iter_rows(min_row=1, max_row=toi_da, values_only=True), start=1):
        chu = [chuan_hoa_ten(v) for v in hang if isinstance(v, str)]
        if all(any(k in c for c in chu) for k in tu_khoa):
            return i
    return None


# ------------------------------------------------------------------ DSR: QTBH
QTBH_COT = [  # (từ khoá tiêu đề đã bỏ dấu, tên trường)
    ("giam gia nhan vien", "giam_gia"),
    ("anh quang tiep khach", "tiep_khach_noi_bo"),
    ("anh quang", "tiep_khach_noi_bo"),
    ("chi chung", "tiep_khach_noi_bo"),
    ("cn anh ha", "cong_no_khach"),
    ("cn khach le", "cong_no_khach"),
    ("cong no nhan vien", "cong_no_khach"),
    ("tru lui", "tru_lui"),
    ("thue gtgt", "vat"),
    ("tong phai thu", "tong_phai_thu"),
    ("ca 1", "ca1"), ("ca 2", "ca2"), ("ck chi", "ck_chu"),
    ("tong cong", "tong_khong_tien_mat"),
    ("tien ban hang can phai nop", "tien_mat"),
    ("ghi chu", "ghi_chu"),
    ("doanh thu", "doanh_thu_gop"),
    ("ngay", "ngay"),
]


def doc_qtbh(wb, den_ngay: date | None = None) -> tuple[list[dict], list[str]]:
    ws = wb["QTBH"]
    hang_td = _tim_hang_tieu_de(ws, ["ngay", "doanh thu"])
    if hang_td is None:
        raise ValueError("Không thấy tiêu đề NGÀY / DOANH THU trong sheet QTBH")
    h1 = next(ws.iter_rows(min_row=hang_td, max_row=hang_td, values_only=True))
    h2 = next(ws.iter_rows(min_row=hang_td + 1, max_row=hang_td + 1, values_only=True))
    vi_tri: dict[int, str] = {}
    for i in range(len(h1)):
        ten = chuan_hoa_ten(str(h2[i] if i < len(h2) and h2[i] else h1[i] or ""))
        for khoa, truong in QTBH_COT:
            if ten.startswith(khoa):
                vi_tri[i] = truong
                break
    ghi_chu_kt: list[str] = []
    ket_qua = []
    for hang in ws.iter_rows(min_row=hang_td + 2, values_only=True):
        gt: dict[str, float] = {}
        ngay = None
        chu_thich = ""
        for i, truong in vi_tri.items():
            v = hang[i] if i < len(hang) else None
            if truong == "ngay":
                ngay = doc_ngay(v)
            elif truong == "ghi_chu":
                chu_thich = str(v or "").strip()
            else:
                so = _so(v)
                if so is not None:
                    gt[truong] = gt.get(truong, 0.0) + so
        if not ngay or not gt.get("doanh_thu_gop"):
            continue
        if den_ngay and ngay > den_ngay:
            continue
        thanh_phan = gt.get("ca1", 0) + gt.get("ca2", 0) + gt.get("ck_chu", 0)
        khong_tm = gt.get("tong_khong_tien_mat")
        if khong_tm is None:
            khong_tm = thanh_phan
            ghi_chu_kt.append(f"{ngay:%d/%m}: ô TỔNG CỘNG (CK/thẻ) lỗi hoặc trống – dùng CA1+CA2+CK = {thanh_phan:,.0f}")
        elif abs(khong_tm - thanh_phan) > 1000:
            ghi_chu_kt.append(f"{ngay:%d/%m}: CA1+CA2+CK chị Quyên = {thanh_phan:,.0f} nhưng ô TỔNG CỘNG = "
                              f"{khong_tm:,.0f} – lấy TỔNG CỘNG (khớp tiền phải nộp), cần soát lại")
            chu_thich = (chu_thich + f" | CA1+CA2+CK={thanh_phan:,.0f}≠TỔNG {khong_tm:,.0f}").strip(" |")
        if "tong_phai_thu" in gt:
            tinh = gt["doanh_thu_gop"] - sum(gt.get(k, 0) for k in ("giam_gia", "tiep_khach_noi_bo",
                                                                     "cong_no_khach", "tru_lui"))
            if abs(tinh - gt["tong_phai_thu"]) > 1000:
                ghi_chu_kt.append(f"{ngay:%d/%m}: TỔNG PHẢI THU trên sổ {gt['tong_phai_thu']:,.0f} ≠ tính lại {tinh:,.0f}")
        ket_qua.append({
            "ngay": ngay.isoformat(), "doanh_thu_gop": gt["doanh_thu_gop"],
            "giam_gia": gt.get("giam_gia", 0), "tiep_khach_noi_bo": gt.get("tiep_khach_noi_bo", 0),
            "tru_lui": gt.get("tru_lui", 0), "cong_no_khach": gt.get("cong_no_khach", 0),
            "chuyen_khoan": khong_tm, "the": 0, "tien_mat": gt.get("tien_mat"),
            "vat": gt.get("vat"), "ghi_chu": chu_thich,
        })
    return ket_qua, ghi_chu_kt


def doc_sales_so_khach(wb) -> dict[int, dict]:
    """Sheet SALES: số khách, số bill theo ngày trong tháng (cột TỔNG DOANH THU)."""
    if "SALES" not in wb.sheetnames:
        return {}
    ws = wb["SALES"]
    for i, hang in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        chu = [chuan_hoa_ten(str(v)) if v is not None else "" for v in hang]
        for j, c in enumerate(chu):
            if c == "tong doanh thu" and j + 1 < len(chu) and chu[j + 1].startswith("so khach"):
                kq = {}
                for h in ws.iter_rows(min_row=i + 1, max_row=i + 40, values_only=True):
                    so_ngay = _so(h[0])
                    if so_ngay and 1 <= so_ngay <= 31:
                        kq[int(so_ngay)] = {"so_khach": _so(h[j + 1]),
                                            "so_hoa_don": _so(h[j + 2]) if chu[j + 2].startswith("tong so bill") else None}
                return kq
    return {}


def doc_pm_ban_mon(wb, thang: date) -> list[dict]:
    """Sheet PM: số lượng từng món theo ngày 1..31."""
    if "PM" not in wb.sheetnames:
        return []
    ws = wb["PM"]
    hang_td = _tim_hang_tieu_de(ws, ["ma hang", "mat hang"])
    if hang_td is None:
        return []
    td = next(ws.iter_rows(min_row=hang_td, max_row=hang_td, values_only=True))
    ten_td = [chuan_hoa_ten(str(v)) if v is not None else "" for v in td]
    i_ma, i_ten, i_gia = ten_td.index("ma hang"), ten_td.index("mat hang"), ten_td.index("don gia")
    cot_ngay = {}
    for j, v in enumerate(td):
        so = v if isinstance(v, (int, float)) else None
        if so and float(so).is_integer() and 1 <= so <= 31 and j > i_gia:
            cot_ngay.setdefault(int(so), j)
    kq = []
    so_ngay = so_ngay_trong_thang(thang)
    for hang in ws.iter_rows(min_row=hang_td + 1, values_only=True):
        ma = hang[i_ma]
        if not ma or not isinstance(ma, str):
            continue
        gia = _so(hang[i_gia]) or 0
        for d, j in cot_ngay.items():
            sl = _so(hang[j], so_luong=True) if j < len(hang) else None
            if sl and d <= so_ngay:
                kq.append({"ngay": thang.replace(day=d).isoformat(), "ma_mon": ma.strip(),
                           "ten_mon": str(hang[i_ten] or "").strip(), "so_luong": sl, "don_gia": gia,
                           "thanh_tien": sl * gia})
    return kq


def nhap_dsr(tep: Path, thu_muc: Path, den_ngay: date | None = None) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(tep, data_only=True)
    bao = []
    dt, ghi_chu = doc_qtbh(wb, den_ngay)
    if dt:
        thang = doc_ngay(dt[0]["ngay"]).replace(day=1)
        khach = doc_sales_so_khach(wb)
        for r in dt:
            k = khach.get(doc_ngay(r["ngay"]).day, {})
            r["so_khach"], r["so_hoa_don"] = k.get("so_khach"), k.get("so_hoa_don")
        n, xoa = ghi_bang(thu_muc, "doanh_thu_ngay.csv", dt, "DSR")
        bao.append(f"QTBH: {n} ngày doanh thu (thay {xoa} dòng cũ), tổng DT gộp {sum(r['doanh_thu_gop'] for r in dt):,.0f}đ")
        bao += [f"  ⚠ {g}" for g in ghi_chu]
        mon = doc_pm_ban_mon(wb, thang)
        if den_ngay:
            mon = [m for m in mon if doc_ngay(m["ngay"]) <= den_ngay]
        if mon:
            n, _ = ghi_bang(thu_muc, "ban_hang_pos.csv", mon, "DSR")
            bao.append(f"PM: {n} dòng bán theo món")
        else:
            bao.append("  ⚠ Sheet PM (bán theo món) trống – chưa tính được xuất kho theo định lượng. "
                       "Cần xuất báo cáo mặt hàng từ iPOS mỗi ngày.")
    if "sổ chi" in wb.sheetnames:
        mh, bo = loc_mua_chiu_trung(thu_muc, doc_so_mua_hang_ws(wb["sổ chi"]), "SMH")
        bao += bo
        if mh:
            n, _ = ghi_bang(thu_muc, "nhap_hang.csv", mh, "SMH")
            bao.append(f"Sổ chi (mua hàng): {n} dòng")
        else:
            bao.append("  ⚠ Sheet 'sổ chi' (sổ mua hàng chi tiết) trống.")
    return bao


# ------------------------------------------------------------------ Sổ mua hàng chi tiết
def doc_so_mua_hang_ws(ws) -> list[dict]:
    hang_td = _tim_hang_tieu_de(ws, ["mat hang", "thanh tien"])
    if hang_td is None:
        return []
    td = [chuan_hoa_ten(str(v)) if v is not None else "" for v in
          next(ws.iter_rows(min_row=hang_td, max_row=hang_td, values_only=True))]

    def tim(*khoa):
        for k in khoa:
            for i, c in enumerate(td):
                if c.startswith(k):
                    return i
        return None

    i_ngay, i_ncc, i_hang = tim("ngay"), tim("ten ncc"), tim("mat hang")
    i_sl, i_gia, i_tien = tim("so luong"), tim("don gia"), tim("thanh tien")
    i_nguoi, i_tm, i_no = tim("nguoi mua"), tim("chi tien mat"), tim("cong no")
    kq = []
    ngay_hien, ncc_hien = None, ""
    for hang in ws.iter_rows(min_row=hang_td + 1, values_only=True):
        n = doc_ngay(hang[i_ngay]) if i_ngay is not None else None
        ngay_hien = n or ngay_hien
        if i_ncc is not None and hang[i_ncc]:
            ncc_hien = str(hang[i_ncc]).strip()
        ten = hang[i_hang] if i_hang is not None else None
        if not ten or not ngay_hien:
            continue
        sl = _so(hang[i_sl], so_luong=True) or 1
        gia = _so(hang[i_gia]) if i_gia is not None else None
        tien = _so(hang[i_tien]) if i_tien is not None else None
        no = (_so(hang[i_no]) or 0) if i_no is not None else 0
        tm = (_so(hang[i_tm]) or 0) if i_tm is not None else 0
        chung = {"ngay": ngay_hien.isoformat(), "ncc": ncc_hien or "Chưa ghi NCC", "ten_hang": str(ten).strip(),
                 "so_luong": sl, "don_gia": gia, "nguoi_mua": str(hang[i_nguoi] or "").strip() if i_nguoi is not None else ""}
        tong = tien if tien is not None else (sl * gia if gia else 0)
        if no and tm and abs(no + tm - tong) < 1:
            kq.append({**chung, "thanh_tien": tm, "hinh_thuc_tt": "tien_mat", "ghi_chu": "tách phần trả ngay"})
            kq.append({**chung, "thanh_tien": no, "hinh_thuc_tt": "cong_no", "ghi_chu": "tách phần nợ"})
        else:
            kq.append({**chung, "thanh_tien": tong, "hinh_thuc_tt": "cong_no" if no else "tien_mat"})
    return kq


def nhap_so_mua_hang(tep: Path, thu_muc: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(tep, data_only=True)
    kq = []
    for ws in wb.worksheets:
        kq += doc_so_mua_hang_ws(ws)
    kq, bo = loc_mua_chiu_trung(thu_muc, kq, "SMH")
    n, xoa = ghi_bang(thu_muc, "nhap_hang.csv", kq, "SMH")
    return bo + [f"Sổ mua hàng: {n} dòng (thay {xoa} dòng cũ), tổng {sum(r['thanh_tien'] for r in kq):,.0f}đ"]


# ------------------------------------------------------------------ Sổ Quỹ tiền mặt
SO_QUY_COT = {  # chữ cái cột -> nhóm chi phí
    "H": "luong", "I": "thoi_vu", "J": "bhxh", "K": "nvl_bep", "L": "nvl_bar", "M": "vpp", "N": "marketing",
    "O": "dien_nuoc", "P": "gas_da", "Q": "tieu_hao", "R": "bao_tri", "S": "decor_tiec", "T": "tiep_khach",
    "U": "do_cung_ship", "V": "ccdc", "W": "mat_bang", "X": "thue_mrc", "Y": "internet", "Z": "rac",
    "AA": "via_he", "AB": "phan_mem",
}
SO_QUY_CONG_NO = {"AC": "bep", "AD": "bar", "AE": "khac"}
# Từ khoá trong diễn giải -> nhóm đúng (sửa lỗi ghi sai cột thường gặp)
PHAN_LOAI_LAI = [
    (r"cam ca|thuc an ca|ho koi|ca koi", "van_hanh_ho_koi"),
    (r"test mon|thu mon|mon moi", "rd_mon_moi"),
]


def _chu_cot(i: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(i + 1)


def nhap_so_quy(tep: Path, thu_muc: Path, thang: date) -> list[str]:
    """Đọc 31 sheet ngày của Sổ Quỹ -> chi_phi.csv (mọi khoản chi tiền mặt, kể cả NVL Bếp/Bar)
    và nhap_hang.csv (các khoản mua chịu ở cột CÔNG NỢ)."""
    from openpyxl import load_workbook

    wb = load_workbook(tep, data_only=True)
    chi_phi, mua_no, bao, dt_so_quy = [], [], [], {}
    for d in range(1, so_ngay_trong_thang(thang) + 1):
        if str(d) not in wb.sheetnames:
            continue
        ngay = thang.replace(day=d)
        ws = wb[str(d)]
        for hang in ws.iter_rows(min_row=11, max_row=60, values_only=True):
            if isinstance(hang[0], str) and chuan_hoa_ten(hang[0]).startswith("cong"):
                break
            gt = {_chu_cot(i): v for i, v in enumerate(hang)}
            dien_giai = str(gt.get("E") or "").strip()
            ct = str(gt.get("D") or gt.get("C") or "").strip()
            ghi = str(gt.get("AG") or "").strip()
            if _so(gt.get("F")):
                dt_so_quy[ngay] = dt_so_quy.get(ngay, 0) + _so(gt["F"])
            for cot, nhom in SO_QUY_COT.items():
                tien = _so(gt.get(cot))
                if not tien:
                    continue
                nhom_dung, ghi_them = nhom, ""
                ten_chuan = chuan_hoa_ten(dien_giai)
                for mau, nhom_moi in PHAN_LOAI_LAI:
                    if re.search(mau, ten_chuan):
                        nhom_dung, ghi_them = nhom_moi, f"phân loại lại từ cột {nhom}"
                ky = ""
                if nhom in ("luong", "thoi_vu", "bhxh"):
                    m = re.search(r"\bt(?:hang)?\s?(\d{1,2})\b", ten_chuan)
                    if m and 1 <= int(m.group(1)) <= 12 and int(m.group(1)) != thang.month:
                        thang_ky = int(m.group(1))
                        nam = thang.year - (1 if thang_ky > thang.month else 0)
                        ky = f"{nam}-{thang_ky:02d}"
                        ghi_them = (ghi_them + f" lương thuộc T{thang_ky}").strip()
                chi_phi.append({"ngay": ngay.isoformat(), "nhom_chi_phi": nhom_dung, "noi_dung": dien_giai,
                                "so_tien": tien, "hinh_thuc_tt": "tien_mat", "ky_phan_bo": ky,
                                "ghi_chu": " ".join(x for x in (ct, ghi, ghi_them) if x)})
            for cot, ma in SO_QUY_CONG_NO.items():
                tien = _so(gt.get(cot))
                if not tien:
                    continue
                m = re.search(r"(\d{1,2})\s*/\s*(\d{1,2})", ghi)
                han = ""
                if m:
                    try:
                        han_d = date(thang.year, int(m.group(2)), int(m.group(1)))
                        if han_d < ngay:
                            han_d = han_d.replace(year=han_d.year + 1)
                        han = han_d.isoformat()
                    except ValueError:
                        han = ""
                mua_no.append({"ngay": ngay.isoformat(), "so_chung_tu": f"{ct}-{ngay:%d%m}", "ncc": dien_giai or "Chưa ghi NCC",
                               "ma_nvl": "", "ten_hang": f"Hàng mua chịu ({ma})",
                               "so_luong": 1, "don_gia": tien, "thanh_tien": tien, "hinh_thuc_tt": "cong_no",
                               "han_thanh_toan": han, "ghi_chu": ghi})
    mua_no, bo = loc_mua_chiu_trung(thu_muc, mua_no, "SQ")
    bao += bo
    n1, x1 = ghi_bang(thu_muc, "chi_phi.csv", chi_phi, "SQ")
    n2, _ = ghi_bang(thu_muc, "nhap_hang.csv", mua_no, "SQ")
    bao.append(f"Sổ quỹ T{thang:%m/%Y}: {n1} khoản chi (thay {x1} dòng cũ), tổng {sum(r['so_tien'] for r in chi_phi):,.0f}đ; "
               f"{n2} khoản mua chịu {sum(r['thanh_tien'] for r in mua_no):,.0f}đ")
    for r in chi_phi:
        if "phân loại lại" in r["ghi_chu"] or "lương thuộc" in r["ghi_chu"]:
            bao.append(f"  ↺ {r['ngay']} {r['noi_dung']} {r['so_tien']:,.0f}đ -> {r['nhom_chi_phi']} {r['ky_phan_bo']} ({r['ghi_chu']})")
    bao += _soat_sheet_tong(tep, sum(r["so_tien"] for r in chi_phi))
    # Đối chiếu doanh thu sổ quỹ với QTBH đã nhập
    tep_dt = Path(thu_muc) / "doanh_thu_ngay.csv"
    if tep_dt.exists() and dt_so_quy:
        with tep_dt.open(encoding="utf-8-sig", newline="") as f:
            qtbh = {doc_ngay(r["ngay"]): doc_so(r["doanh_thu_gop"]) for r in csv.DictReader(f)}
        for n, v in sorted(dt_so_quy.items()):
            if n in qtbh and abs(qtbh[n] - v) > 1000:
                bao.append(f"  ⚠ {n:%d/%m}: doanh thu sổ quỹ {v:,.0f} ≠ QTBH {qtbh[n]:,.0f} (lệch {v - qtbh[n]:,.0f})")
    return bao


def _soat_sheet_tong(tep: Path, tong_da_nhap: float) -> list[str]:
    """Sheet TỔNG hay có khoản chi gõ thẳng vào ô công thức (không nằm ở sheet ngày nào).
    Liệt kê ra để ghi nhận đúng (trả nợ NCC, thuế...) thay vì bỏ sót."""
    from openpyxl import load_workbook

    wb_ct = load_workbook(tep, data_only=False)
    wb_gt = load_workbook(tep, data_only=True)
    if "TỔNG" not in wb_gt.sheetnames:
        return []
    ra = []
    ws_gt, ws_ct = wb_gt["TỔNG"], wb_ct["TỔNG"]
    for hang_gt, hang_ct in zip(ws_gt.iter_rows(min_row=1, max_row=80), ws_ct.iter_rows(min_row=1, max_row=80)):
        for o_gt, o_ct in zip(hang_gt, hang_ct):
            ct = o_ct.value
            if isinstance(ct, str) and ct.startswith("=") and re.fullmatch(r"=[\d+\-.\s]+", ct):
                so = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", ct)]
                ra.append(f"  ⚠ Sheet TỔNG ô {o_ct.coordinate} có {len(so)} khoản chi gõ tay KHÔNG nằm ở sheet ngày nào "
                          f"(tổng {sum(so):,.0f}đ): {', '.join(f'{x:,.0f}' for x in so)}. "
                          "Cần ghi nhận vào thanh_toan_ncc.csv / chi_phi.csv với đúng ngày & nội dung.")
    return ra


# ------------------------------------------------------------------ iPOS: bán theo mặt hàng (1 ngày)
def nhap_ipos_mat_hang(tep: Path, thu_muc: Path, ngay: date) -> list[str]:
    tep = Path(tep)
    if tep.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(tep, data_only=True).active
        hang = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with tep.open(encoding="utf-8-sig", newline="") as f:
            hang = list(csv.reader(f))
    for i, h in enumerate(hang[:30]):
        td = [chuan_hoa_ten(str(v or "")) for v in h]
        i_ma = next((j for j, c in enumerate(td) if c in ("ma hang", "ma mon", "ma", "item code", "ma mat hang")), None)
        i_ten = next((j for j, c in enumerate(td) if c in ("ten hang", "ten mon", "mat hang", "ten mat hang", "item")), None)
        i_sl = next((j for j, c in enumerate(td) if c in ("so luong", "sl", "qty", "so luong ban")), None)
        i_tien = next((j for j, c in enumerate(td) if c.startswith(("thanh tien", "doanh thu", "tong tien"))), None)
        if i_sl is not None and (i_ma is not None or i_ten is not None):
            break
    else:
        raise ValueError("Không nhận ra cột Mã/Tên hàng và Số lượng trong file iPOS")
    kq = []
    for h in hang[i + 1:]:
        ma = str(h[i_ma] or "").strip() if i_ma is not None else ""
        ten = str(h[i_ten] or "").strip() if i_ten is not None else ""
        sl = _so(h[i_sl], so_luong=True)
        if not sl or not (ma or ten) or chuan_hoa_ten(ma or ten).startswith("tong"):
            continue
        tien = _so(h[i_tien]) if i_tien is not None else None
        kq.append({"ngay": ngay.isoformat(), "ma_mon": ma or ten, "ten_mon": ten, "so_luong": sl,
                   "don_gia": (tien / sl) if tien else None, "thanh_tien": tien})
    n, xoa = ghi_bang(thu_muc, "ban_hang_pos.csv", kq, "IPOS")
    return [f"iPOS {ngay:%d/%m}: {n} món (thay {xoa} dòng cũ), tổng {sum(r['thanh_tien'] or 0 for r in kq):,.0f}đ"]
