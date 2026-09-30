"""Lập báo cáo ngày / tháng: Markdown (đọc nhanh, gửi Zalo/email) + Excel (lưu trữ, đối chiếu)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from .cong_no import (HINH_THUC, chuan_hinh_thuc, kiem_soat_hoa_don, tinh_cong_no_khach,
                      tinh_cong_no_ncc, tong_hop_mua_hang)
from .doanh_thu import doanh_thu_ngay, kiem_soat_doanh_thu, top_mon
from .du_lieu import NHOM_CHI_PHI, DuLieu, trong_khoang
from .lai_lo import dong_tien, lai_lo
from .nhap_xuat_ton import tinh_nhap_xuat_ton, tong_hop_theo_nhom, tong_hop_theo_nvl
from .tien_ich import cac_ngay, dau_thang, dinh_dang_pct, dinh_dang_so, dinh_dang_tien, thu_trong_tuan

MUC = {"cao": "🔴", "trung_binh": "🟠", "thap": "🟡"}
THU_TU_MUC = {"cao": 0, "trung_binh": 1, "thap": 2}
TEN_THU = {"T2": "Thứ Hai", "T3": "Thứ Ba", "T4": "Thứ Tư", "T5": "Thứ Năm", "T6": "Thứ Sáu",
           "T7": "Thứ Bảy", "CN": "Chủ Nhật"}


def t(so) -> str:
    return dinh_dang_tien(so) if so is not None else "–"


def p(ty_le) -> str:
    return dinh_dang_pct(ty_le) if ty_le is not None else "–"


def _bang(tieu_de: list[str], dong: list[list], can_phai: set[int] | None = None) -> str:
    can_phai = can_phai if can_phai is not None else set(range(1, len(tieu_de)))
    ke = ["---:" if i in can_phai else "---" for i in range(len(tieu_de))]
    ra = ["| " + " | ".join(tieu_de) + " |", "| " + " | ".join(ke) + " |"]
    ra += ["| " + " | ".join(str(o) for o in d) + " |" for d in dong]
    return "\n".join(ra)


def do_day_du(dl: DuLieu, tu: date, den: date) -> dict:
    """Ngày nào có doanh thu nhưng chưa có chi phí/sổ quỹ -> lãi lỗ ngày đó chưa đủ tin cậy."""
    co_dt = {r["ngay"] for r in trong_khoang(dl["doanh_thu_ngay.csv"], tu, den)}
    co_chi = {r["ngay"] for r in trong_khoang(dl["chi_phi.csv"], tu, den)}
    tong_ngay = list(cac_ngay(tu, den))
    return {"so_ngay": len(tong_ngay),
            "thieu_doanh_thu": [n for n in tong_ngay if n not in co_dt],
            "thieu_chi_phi": [n for n in tong_ngay if n in co_dt and n not in co_chi]}


def _ds_ngay(ds: list[date]) -> str:
    return ", ".join(f"{n:%d/%m}" for n in ds[:10]) + (f" … (+{len(ds) - 10} ngày)" if len(ds) > 10 else "")


def tinh_toan(dl: DuLieu, tu: date, den: date) -> dict:
    """Chạy toàn bộ nghiệp vụ cho kỳ [tu, den] và lũy kế từ đầu tháng của `den`."""
    dau = dau_thang(den)
    bat_dau = min(tu, dau)
    kho = tinh_nhap_xuat_ton(dl, bat_dau, den)
    ky = lai_lo(dl, kho, tu, den)
    luy_ke = lai_lo(dl, kho, dau, den) if (tu, den) != (dau, den) else ky
    cong_no = tinh_cong_no_ncc(dl, den)
    canh_bao = (
        [c for c in kho.canh_bao if tu <= c["ngay"] <= den or c["loai"] in ("Hao hụt", "Định lượng")]
        + kiem_soat_hoa_don(dl, tu, den)
        + kiem_soat_doanh_thu(dl, tu, den)
        + cong_no.canh_bao
        + [{"muc": "cao", "loai": "Dữ liệu", "ngay": den, "noi_dung": l} for l in dl.loi_nap]
    )
    if ky["chi_loai_khoi_ky"]:
        tong = sum(r["so_tien"] for r in ky["chi_loai_khoi_ky"])
        canh_bao.append({"muc": "thap", "loai": "Kỳ chi phí", "ngay": den,
                         "noi_dung": f"{len(ky['chi_loai_khoi_ky'])} khoản chi {t(tong)}đ trả trong kỳ nhưng "
                                     "thuộc kỳ khác (cột ky_phan_bo) – đã loại khỏi lãi lỗ kỳ này"})
    du = do_day_du(dl, dau_thang(den) if tu == den else tu, den)
    if du["thieu_chi_phi"]:
        canh_bao.append({"muc": "cao", "loai": "THIẾU SỐ LIỆU", "ngay": den,
                         "noi_dung": f"{len(du['thieu_chi_phi'])} ngày có doanh thu nhưng CHƯA có Sổ Quỹ/chi phí "
                                     f"({_ds_ngay(du['thieu_chi_phi'])}) -> giá vốn & chi phí thiếu, lợi nhuận đang bị "
                                     "THỔI PHỒNG. Cần nhập Sổ Quỹ các ngày này."})
    canh_bao.sort(key=lambda c: (THU_TU_MUC.get(c["muc"], 9), c["ngay"]))
    return {
        "tu": tu, "den": den, "kho": kho, "ky": ky, "luy_ke": luy_ke, "cong_no": cong_no,
        "cong_no_khach": tinh_cong_no_khach(dl, den), "mua_hang": tong_hop_mua_hang(dl, tu, den),
        "dong_tien": dong_tien(dl, tu, den), "canh_bao": canh_bao, "top_mon": top_mon(dl, tu, den),
        "ten": dl.cau_hinh.get("ten_nha_hang", "RUBY KOI BISTRO"), "dl": dl, "do_day_du": du,
    }


# ---------------------------------------------------------------- Markdown
def _muc_kpi(kq: dict, la_ngay: bool) -> str:
    ky, lk = kq["ky"], kq["luy_ke"]
    dt_ky, dt_lk = ky["doanh_thu"], lk["doanh_thu"]
    cot = ["Chỉ tiêu", "Hôm nay" if la_ngay else "Cả kỳ"] + (["Lũy kế tháng"] if la_ngay else [])

    thieu = kq["do_day_du"]["thieu_chi_phi"]
    thieu_ky = any(kq["tu"] <= n <= kq["den"] for n in thieu)
    thieu_lk = bool(thieu)
    CHUA = "⚠ chưa đủ số liệu"

    def dong(ten, a, b, la_chi_phi=False):
        if la_chi_phi:
            a = CHUA if thieu_ky else a
            b = CHUA if thieu_lk else b
        return [ten, a] + ([b] if la_ngay else [])

    ds = [
        dong("Doanh thu gộp", t(ky["doanh_thu_gop"]), t(lk["doanh_thu_gop"])),
        dong("Doanh thu thuần", f"**{t(ky['doanh_thu_thuan'])}**", f"**{t(lk['doanh_thu_thuan'])}**"),
        dong("Số khách / TB đầu người", f"{t(dt_ky['cong'].get('so_khach'))} / {t(dt_ky['tb_khach'])}",
             f"{t(dt_lk['cong'].get('so_khach'))} / {t(dt_lk['tb_khach'])}"),
        dong("Giá vốn (COGS)", t(ky["gia_von"]["tong"]), t(lk["gia_von"]["tong"]), la_chi_phi=True),
        dong("Food cost %", p(ky["ty_le"]["food_cost"]), p(lk["ty_le"]["food_cost"]), la_chi_phi=True),
        dong("Hao hụt (huỷ + thất thoát)", t(ky["gia_von"].get("huy", 0) + ky["gia_von"].get("that_thoat", 0)),
             t(lk["gia_von"].get("huy", 0) + lk["gia_von"].get("that_thoat", 0))),
        dong("Lợi nhuận gộp", t(ky["loi_nhuan_gop"]), t(lk["loi_nhuan_gop"]), la_chi_phi=True),
        dong("EBITDA (trước khấu hao)", f"**{t(ky['ebitda'])}**", f"**{t(lk['ebitda'])}**", la_chi_phi=True),
        dong("Lợi nhuận kế toán", t(ky["loi_nhuan"]), t(lk["loi_nhuan"]), la_chi_phi=True),
        dong("Chênh lệch quỹ tiền mặt", t(dt_ky["cong"].get("chenh_lech_quy")),
             t(dt_lk["cong"].get("chenh_lech_quy"))),
    ]
    ra = _bang(cot, ds)
    ra += (f"\n\nCông nợ NCC còn phải trả: **{t(kq['cong_no'].tong_con_no())}đ** · "
           f"Khách còn nợ: **{t(kq['cong_no_khach']['con_phai_thu'])}đ**")
    if dt_lk["hoa_von"]:
        ra += (f" · Lũy kế đạt **{p(dt_lk['dat_hoa_von'])}** điểm hoà vốn tháng ({t(dt_lk['hoa_von'])}đ), "
               f"dự báo cả tháng ≈ {t(dt_lk['du_bao_thang'])}đ")
    return ra


def _muc_doanh_thu(kq: dict, la_ngay: bool) -> str:
    dl = kq["dl"]
    ra = []
    if la_ngay:
        d = doanh_thu_ngay(dl, kq["den"])
        if d is None:
            return "_Chưa có số QTBH ngày này._"
        ra.append(_bang(["Khoản", "Số tiền"], [
            ["Doanh thu gộp (POS)", t(d["doanh_thu_gop"])],
            ["(−) Giảm giá / voucher / chiết khấu", t(d["giam_gia"])],
            ["(−) Tiếp khách, chi chung, nội bộ", t(d["tiep_khach_noi_bo"])],
            ["(−) Trừ lui", t(d["tru_lui"])],
            ["**= Doanh thu thuần**", f"**{t(d['doanh_thu_thuan'])}**"],
            ["(−) Khách nợ (chuyển sang phải thu)", t(d["cong_no_khach"])],
            ["= Thực thu trong ngày", t(d["thuc_thu"])],
            ["   · Chuyển khoản", t(d["chuyen_khoan"])],
            ["   · Thẻ / VISA", t(d["the"])],
            ["   · Tiền mặt phải nộp", t(d["tien_mat"])],
            ["(−) Chi tiền mặt từ quỹ thu ngân", t(d["chi_tien_mat"]["tong_chi"])],
            ["(+) Thu nợ khách bằng tiền mặt", t(d["chi_tien_mat"]["thu_no_khach"])],
            ["= Tiền mặt phải còn trong két", t(d["tien_mat_phai_con"])],
            ["Tiền mặt đếm thực tế", t(d["tien_mat_dem"])],
            ["**Chênh lệch quỹ**", f"**{t(d['chenh_lech_quy'])}**"],
        ]))
        if d["pos_tong_mon"] is not None:
            ra.append(f"\nĐối chiếu: tổng bán theo món POS = {t(d['pos_tong_mon'])}đ "
                      f"(lệch QTBH {t(d['pos_tong_mon'] - d['doanh_thu_gop'])}đ)")
    else:
        dt = kq["ky"]["doanh_thu"]
        c = dt["cong"]
        ra.append(_bang(["Khoản", "Số tiền"], [
            ["Doanh thu gộp", t(c.get("doanh_thu_gop"))], ["(−) Giảm trừ", t(c.get("giam_tru"))],
            ["**Doanh thu thuần**", f"**{t(c.get('doanh_thu_thuan'))}**"],
            ["Khách nợ phát sinh", t(c.get("cong_no_khach"))], ["Thực thu", t(c.get("thuc_thu"))],
            ["Tổng chênh lệch quỹ", t(c.get("chenh_lech_quy"))],
            ["Số ngày có số liệu", dt["so_ngay_co_so"]],
        ]))
        ra.append("\n**Theo thứ trong tuần**\n")
        ra.append(_bang(["Thứ", "Số ngày", "DT thuần", "BQ/ngày", "Tỷ trọng"],
                        [[k, v["so_ngay"], t(v["doanh_thu_thuan"]), t(v["binh_quan"]), p(v["ty_trong"])]
                         for k, v in dt["theo_thu"].items() if v["so_ngay"]]))
        ra.append("\n**Cơ cấu thanh toán:** " + " · ".join(
            f"{k} {p(v)}" for k, v in dt["ty_trong_thanh_toan"].items()))
        ra.append("\n**Bảng từng ngày**\n")
        ra.append(_bang(["Ngày", "Thứ", "DT gộp", "Giảm trừ", "DT thuần", "Khách nợ", "CK+Thẻ", "Tiền mặt",
                         "Lệch quỹ"],
                        [[f"{d['ngay']:%d/%m}", d["thu"], t(d["doanh_thu_gop"]), t(d["giam_tru"]),
                          t(d["doanh_thu_thuan"]), t(d["cong_no_khach"]), t(d["chuyen_khoan"] + d["the"]),
                          t(d["tien_mat"]), t(d["chenh_lech_quy"])] for d in dt["ngay"]]))
    if kq["top_mon"]:
        ra.append("\n**Món bán chạy (theo doanh thu)**\n")
        ra.append(_bang(["Món", "SL", "Doanh thu"],
                        [[m["ten_mon"] or m["ma_mon"], dinh_dang_so(m["so_luong"]), t(m["thanh_tien"])]
                         for m in kq["top_mon"]]))
    return "\n".join(ra)


def _muc_kho(kq: dict, la_ngay: bool) -> str:
    kho = kq["kho"]
    ds = [d for d in kho.dong if kq["tu"] <= d.ngay <= kq["den"]]
    tong = tong_hop_theo_nvl(ds)
    ty_le_lk = {k: v["ty_le_hao_hut"] for k, v in tong_hop_theo_nvl(kho.dong).items()}
    if not tong:
        return "_Chưa có danh mục NVL theo dõi tồn (danh_muc_nvl.csv, quan_ly_ton = co)._"
    ra = [_bang(
        ["NVL", "ĐVT", "Tồn đầu", "Nhập", "Xuất bán", "Huỷ", "Thất thoát", "Thừa", "Tồn cuối",
         "Kiểm kê", "Giá trị tồn", "Hao hụt % lũy kế" if la_ngay else "Hao hụt %"],
        [[v["ten_nvl"], v["don_vi"], dinh_dang_so(v["ton_dau"]), dinh_dang_so(v["nhap"]),
          dinh_dang_so(v["xuat_ban"]), dinh_dang_so(v["huy"]), dinh_dang_so(v["that_thoat"]),
          dinh_dang_so(v["thua"]), dinh_dang_so(v["ton_cuoi"]),
          f"{v['so_lan_kiem']} lần" if v["so_lan_kiem"] else "chưa", t(v["gia_tri_ton_cuoi"]),
          p(ty_le_lk.get(v["ma_nvl"]) if la_ngay else v["ty_le_hao_hut"])]
         for v in sorted(tong.values(), key=lambda x: (x["nhom"], x["ten_nvl"]))
         if not la_ngay or any((v["nhap"], v["xuat_ban"], v["huy"], v["that_thoat"], v["thua"]))],
        can_phai=set(range(2, 12)))]
    nhom = tong_hop_theo_nhom(kho.dong if la_ngay else ds)
    ra.append(f"\n**Theo nhóm – {'lũy kế tháng' if la_ngay else 'cả kỳ'} (giá trị, đồng)**\n")
    ra.append(_bang(["Nhóm", "Tồn đầu", "Nhập", "Xuất bán", "Huỷ", "Thất thoát", "Tồn cuối", "Giá vốn thực tế",
                     "Hao hụt %"],
                    [[k, t(g["gia_tri_ton_dau"]), t(g["gia_tri_nhap"]), t(g["gia_tri_xuat_ban"]),
                      t(g["gia_tri_huy"]), t(g["gia_tri_that_thoat"]), t(g["gia_tri_ton_cuoi"]),
                      t(g["gia_von_thuc_te"]), p(g["ty_le_hao_hut"])] for k, g in sorted(nhom.items())]))
    gv = kq["ky"]["gia_von"]
    ra.append(f"\nGiá vốn kỳ = mua NVL (Sổ Quỹ {t(gv.get('mua_so_quy', 0))} + mua chịu/sổ nhập "
              f"{t(gv.get('mua_nhap_hang', 0))}) + giảm tồn kho {t(gv.get('giam_ton', 0))} = **{t(gv['tong'])}đ**. "
              f"Theo định lượng (lý thuyết) hàng tồn kho: {t(gv.get('gia_von_ly_thuyet', 0))}đ")
    return "\n".join(ra)


def _muc_mua_hang(kq: dict, la_ngay: bool) -> str:
    dl, mh, cn = kq["dl"], kq["mua_hang"], kq["cong_no"]
    ra = [f"Sổ mua hàng/chi chi tiết: {t(mh['tong'])}đ ({mh['so_dong']} dòng): " + " · ".join(
        f"{HINH_THUC.get(k, k)} {t(v)}" for k, v in mh["theo_hinh_thuc"].items()) +
        f" · Trả nợ NCC trong kỳ: {t(mh['tra_no_ncc'])}đ"]
    if la_ngay:
        dong = trong_khoang(dl["nhap_hang.csv"], kq["tu"], kq["den"])
        if dong:
            ra.append("\n" + _bang(["NCC", "Chứng từ", "Mặt hàng", "SL", "Đơn giá", "Thành tiền", "Thanh toán"],
                                   [[r["ncc"], r.get("so_chung_tu", ""), r["ten_hang"], dinh_dang_so(r["so_luong"]),
                                     t(r["don_gia"]), t(r["thanh_tien"]),
                                     HINH_THUC[chuan_hinh_thuc(r.get("hinh_thuc_tt", ""))]] for r in dong],
                                   can_phai={3, 4, 5}))
    elif mh["theo_ncc"]:
        ra.append("\n" + _bang(["NCC", "Tiền mặt", "Chuyển khoản", "Công nợ", "Tổng"],
                               [[k, t(v.get("tien_mat", 0)), t(v.get("chuyen_khoan", 0)), t(v.get("cong_no", 0)),
                                 t(v["tong"])] for k, v in sorted(mh["theo_ncc"].items(), key=lambda x: -x[1]["tong"])[:20]]))
        if len(mh["theo_ncc"]) > 20:
            ra.append(f"_… và {len(mh['theo_ncc']) - 20} người bán khác (xem file Excel, sheet Mua hàng)._")
    tuoi = cn.tuoi_no()
    ra.append(f"\n**Công nợ nhà cung cấp tại {kq['den']:%d/%m/%Y}: {t(cn.tong_con_no())}đ**\n")
    if tuoi:
        ra.append(_bang(["NCC", "Chưa đến hạn", "Quá hạn 1–7 ngày", "8–30 ngày", "> 30 ngày", "Tổng"],
                        [[k, t(v.get("chua_den_han", 0)), t(v.get("qua_han_1_7", 0)), t(v.get("qua_han_8_30", 0)),
                          t(v.get("qua_han_tren_30", 0)), t(v["tong"])]
                         for k, v in sorted(tuoi.items(), key=lambda x: -x[1]["tong"])]))
        ra.append("\n" + _bang(["NCC", "Chứng từ", "Ngày", "Hạn", "Số tiền", "Đã trả", "Còn nợ"],
                               [[h.ncc, h.so_chung_tu + (" (đầu kỳ)" if h.dau_ky else ""), f"{h.ngay:%d/%m}",
                                 f"{h.han:%d/%m}", t(h.so_tien), t(h.da_tra), t(h.con_lai)]
                                for h in cn.con_no()], can_phai={4, 5, 6}))
    else:
        ra.append("Không còn nợ nhà cung cấp.")
    kh = kq["cong_no_khach"]
    ra.append(f"\n**Công nợ khách hàng:** phát sinh {t(kh['phat_sinh'])} − đã thu {t(kh['da_thu'])} = "
              f"**{t(kh['con_phai_thu'])}đ còn phải thu**")
    if kh["theo_khach"]:
        ra.append(_bang(["Khách", "Còn nợ"], [[k, t(v)] for k, v in kh["theo_khach"].items()]))
    return "\n".join(ra)


def _muc_lai_lo(kq: dict, la_ngay: bool) -> str:
    cot_ky = "Hôm nay" if la_ngay else "Cả kỳ"
    ky, lk = kq["ky"], kq["luy_ke"]
    hai = la_ngay

    def dong(ten, a, b, pa=None, pb=None):
        d = [ten, a, p(pa) if pa is not None else ""]
        return d + ([b, p(pb) if pb is not None else ""] if hai else [])

    dtt_k, dtt_l = ky["doanh_thu_thuan"], lk["doanh_thu_thuan"]
    r = lambda x, m: (x / m) if m else None  # noqa: E731
    ds = [dong("Doanh thu gộp", t(ky["doanh_thu_gop"]), t(lk["doanh_thu_gop"])),
          dong("(−) Giảm trừ doanh thu", t(ky["giam_tru"]), t(lk["giam_tru"])),
          dong("**Doanh thu thuần**", f"**{t(dtt_k)}**", f"**{t(dtt_l)}**", 1 if dtt_k else None, 1 if dtt_l else None),
          dong("(−) Giá vốn NVL", t(ky["gia_von"]["tong"]), t(lk["gia_von"]["tong"]),
               ky["ty_le"]["food_cost"], lk["ty_le"]["food_cost"]),
          dong("**Lợi nhuận gộp**", f"**{t(ky['loi_nhuan_gop'])}**", f"**{t(lk['loi_nhuan_gop'])}**",
               r(ky["loi_nhuan_gop"], dtt_k), r(lk["loi_nhuan_gop"], dtt_l))]
    for nhom in sorted(set(ky["nhan_su"]) | set(lk["nhan_su"])):
        ds.append(dong(f"(−) {NHOM_CHI_PHI.get(nhom, nhom)}", t(ky["nhan_su"].get(nhom, 0)),
                       t(lk["nhan_su"].get(nhom, 0))))
    ds.append(dong("= Chi phí nhân sự", t(ky["tong_nhan_su"]), t(lk["tong_nhan_su"]),
                   ky["ty_le"]["nhan_su"], lk["ty_le"]["nhan_su"]))
    for nhom in sorted(set(ky["van_hanh"]) | set(lk["van_hanh"])):
        ds.append(dong(f"(−) {NHOM_CHI_PHI.get(nhom, nhom)}", t(ky["van_hanh"].get(nhom, 0)),
                       t(lk["van_hanh"].get(nhom, 0))))
    ds += [dong("= Chi phí vận hành", t(ky["tong_van_hanh"]), t(lk["tong_van_hanh"]),
                ky["ty_le"]["van_hanh"], lk["ty_le"]["van_hanh"]),
           dong("**EBITDA**", f"**{t(ky['ebitda'])}**", f"**{t(lk['ebitda'])}**",
                ky["ty_le"]["ebitda"], lk["ty_le"]["ebitda"]),
           dong("(−) Khấu hao (phân bổ)", t(ky["khau_hao"]), t(lk["khau_hao"])),
           dong("(−) Thuế nộp & lãi vay", t(ky["lai_vay"]), t(lk["lai_vay"])),
           dong("**Lợi nhuận kế toán**", f"**{t(ky['loi_nhuan'])}**", f"**{t(lk['loi_nhuan'])}**",
                ky["ty_le"]["loi_nhuan"], lk["ty_le"]["loi_nhuan"])]
    cot = ["Chỉ tiêu", cot_ky, "% DTT"] + (["Lũy kế tháng", "% DTT"] if hai else [])
    ra = [_bang(cot, ds)]
    ra.append(f"\nPrime cost (NVL + nhân sự): {p(ky['ty_le']['prime_cost'])}"
              + (f" hôm nay · {p(lk['ty_le']['prime_cost'])} lũy kế" if hai else ""))
    dtien = kq["dong_tien"]
    ra.append(f"\n**Dòng tiền {'hôm nay' if la_ngay else 'cả kỳ'}:** thu {t(dtien['tong_thu'])} − chi "
              f"{t(dtien['tong_chi'])} (chi phí {t(sum(dtien['chi_phi'].values()))}, mua trả ngay "
              f"{t(dtien['mua_hang_tra_ngay'])}, trả nợ NCC {t(dtien['tra_no_ncc'])}) = "
              f"**{t(dtien['dong_tien_thuan'])}đ**. Nghĩa vụ trả ngân hàng/tháng (chưa trừ): "
              f"{t(dtien['tra_ngan_hang_thang'])}đ")
    return "\n".join(ra)


def _muc_canh_bao(kq: dict) -> str:
    if not kq["canh_bao"]:
        return "✅ Không có cảnh báo. Số liệu khớp."
    return "\n".join(f"{i}. {MUC.get(c['muc'], '')} **[{c['loai']}]** {c['ngay']:%d/%m} – {c['noi_dung']}"
                     for i, c in enumerate(kq["canh_bao"], 1))


def _banner_thieu(kq: dict) -> str:
    du = kq["do_day_du"]
    dong = []
    if du["thieu_chi_phi"]:
        dong.append(f"> ⚠️ **SỐ LIỆU CHƯA ĐỦ – LỢI NHUẬN CHƯA DÙNG ĐƯỢC:** {len(du['thieu_chi_phi'])}/{du['so_ngay']} ngày "
                    f"trong kỳ lũy kế chưa có Sổ Quỹ/chi phí ({_ds_ngay(du['thieu_chi_phi'])}). Doanh thu chính xác; "
                    "giá vốn, chi phí và lợi nhuận sẽ đúng khi nhập đủ Sổ Quỹ.")
    if du["thieu_doanh_thu"]:
        dong.append(f"> ⚠️ Chưa có QTBH/POS các ngày: {_ds_ngay(du['thieu_doanh_thu'])}.")
    return ("\n".join(dong) + "\n") if dong else ""


def ra_markdown(kq: dict, loai: str) -> str:
    la_ngay = loai == "ngay"
    den = kq["den"]
    if la_ngay:
        tieu_de = f"BÁO CÁO KẾ TOÁN NỘI BỘ NGÀY {den:%d/%m/%Y} ({TEN_THU[thu_trong_tuan(den)]})"
    else:
        tieu_de = f"BÁO CÁO KẾ TOÁN NỘI BỘ THÁNG {den:%m/%Y} ({kq['tu']:%d/%m} – {den:%d/%m})"
    so_cao = sum(1 for c in kq["canh_bao"] if c["muc"] == "cao")
    phan = [
        f"# {kq['ten']}\n## {tieu_de}\n",
        f"Đơn vị: VNĐ · Cảnh báo: {so_cao} mức cao / {len(kq['canh_bao'])} tổng · "
        "Số liệu lấy từ QTBH/POS, sổ nhập hàng, sổ quỹ, kiểm kho.\n",
        _banner_thieu(kq),
        "### 1. Tóm tắt nhanh\n", _muc_kpi(kq, la_ngay),
        "\n### 2. Doanh thu & đối soát quỹ tiền mặt\n", _muc_doanh_thu(kq, la_ngay),
        "\n### 3. Nhập – Xuất – Tồn – Hao hụt\n", _muc_kho(kq, la_ngay),
        "\n### 4. Hoá đơn mua hàng & công nợ\n", _muc_mua_hang(kq, la_ngay),
        "\n### 5. Kết quả kinh doanh (P&L) & dòng tiền\n", _muc_lai_lo(kq, la_ngay),
        "\n### 6. Cảnh báo & việc cần xử lý\n", _muc_canh_bao(kq),
        "\n---\n_Phương pháp: giá vốn = tồn đầu + nhập − tồn cuối (bình quân gia quyền); mặt bằng, BHXH, "
        "khấu hao phân bổ đều theo ngày; khách nợ không trừ khỏi doanh thu mà ghi phải thu._",
    ]
    return "\n".join(phan) + "\n"


# ---------------------------------------------------------------- Excel
def ra_excel(kq: dict, loai: str, tep: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    dam = Font(bold=True)
    nen = PatternFill("solid", fgColor="8B1A1A")
    chu_trang = Font(bold=True, color="FFFFFF")
    so = '#,##0;[Red](#,##0);"-"'
    so_le = '#,##0.00;[Red](#,##0.00);"-"'
    pct = "0.0%"

    def sheet(ten, tieu_de, cot, dong, dinh_dang=None):
        ws = wb.create_sheet(ten)
        ws.append([kq["ten"]])
        ws.append([tieu_de])
        ws["A1"].font = ws["A2"].font = dam
        ws.append([])
        ws.append(cot)
        for o in ws[4]:
            o.font, o.fill = chu_trang, nen
            o.alignment = Alignment(wrap_text=True, vertical="center")
        for d in dong:
            ws.append(d)
        for i, _ in enumerate(cot, 1):
            ws.column_dimensions[get_column_letter(i)].width = 14 if i > 1 else 34
            fmt = (dinh_dang or {}).get(i, so)
            for hang in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                for o in hang:
                    if isinstance(o.value, (int, float)):
                        o.number_format = fmt
        ws.freeze_panes = "B5"
        return ws

    wb.remove(wb.active)
    ky, lk, den = kq["ky"], kq["luy_ke"], kq["den"]
    ten_ky = f"{den:%d/%m/%Y}" if loai == "ngay" else f"T{den:%m/%Y}"

    def pl_dong(x):
        g = x["gia_von"]
        ds = [["Doanh thu gộp", x["doanh_thu_gop"]], ["Giảm trừ doanh thu", x["giam_tru"]],
              ["DOANH THU THUẦN", x["doanh_thu_thuan"]],
              ["Mua NVL – Sổ Quỹ (Bếp + Bar)", g.get("mua_so_quy", 0)],
              ["Mua NVL – mua chịu / sổ nhập hàng", g.get("mua_nhap_hang", 0)],
              ["(+) Giảm tồn kho (tồn đầu − tồn cuối)", g.get("giam_ton", 0)],
              ["TỔNG GIÁ VỐN", g["tong"]], ["  trong đó huỷ", g.get("huy", 0)],
              ["  trong đó thất thoát", g.get("that_thoat", 0)], ["LỢI NHUẬN GỘP", x["loi_nhuan_gop"]]]
        ds += [[f"Nhân sự – {NHOM_CHI_PHI.get(k, k)}", v] for k, v in sorted(x["nhan_su"].items())]
        ds += [[f"Vận hành – {NHOM_CHI_PHI.get(k, k)}", v] for k, v in sorted(x["van_hanh"].items())]
        ds += [["EBITDA", x["ebitda"]], ["Khấu hao", x["khau_hao"]], ["Thuế nộp & lãi vay", x["lai_vay"]],
               ["LỢI NHUẬN KẾ TOÁN", x["loi_nhuan"]]]
        return ds

    ds_ky = pl_dong(ky)
    if loai == "ngay":
        luy = dict((a, b) for a, b in pl_dong(lk))
        ten_dong = [a for a, _ in ds_ky] + [a for a, _ in pl_dong(lk) if a not in dict(ds_ky)]
        ky_d = dict(ds_ky)
        dtt, dtl = ky["doanh_thu_thuan"], lk["doanh_thu_thuan"]
        dong = [[a, ky_d.get(a, 0), (ky_d.get(a, 0) / dtt) if dtt else None, luy.get(a, 0),
                 (luy.get(a, 0) / dtl) if dtl else None] for a in ten_dong]
        ws = sheet("P&L", f"KẾT QUẢ KINH DOANH – {ten_ky}", ["Chỉ tiêu", "Hôm nay", "% DTT", "Lũy kế tháng", "% DTT"],
                   dong, {3: pct, 5: pct})
    else:
        dtt = ky["doanh_thu_thuan"]
        ws = sheet("P&L", f"KẾT QUẢ KINH DOANH – {ten_ky}", ["Chỉ tiêu", "Số tiền", "% DTT"],
                   [[a, b, (b / dtt) if dtt else None] for a, b in ds_ky], {3: pct})
    for hang in ws.iter_rows(min_row=5):
        if hang[0].value and hang[0].value.isupper():
            for o in hang:
                o.font = dam

    dt = ky["doanh_thu"]
    sheet("Doanh thu", f"DOANH THU & ĐỐI SOÁT QUỸ – {ten_ky}",
          ["Ngày", "Thứ", "DT gộp", "Giảm giá", "Tiếp khách/nội bộ", "Trừ lui", "DT thuần", "Khách nợ",
           "Thực thu", "Chuyển khoản", "Thẻ", "Tiền mặt", "Chi TM từ quỹ", "TM phải còn", "TM đếm", "Lệch quỹ",
           "Số khách", "TB/khách", "POS theo món"],
          [[f"{d['ngay']:%d/%m/%Y}", d["thu"], d["doanh_thu_gop"], d["giam_gia"], d["tiep_khach_noi_bo"],
            d["tru_lui"], d["doanh_thu_thuan"], d["cong_no_khach"], d["thuc_thu"], d["chuyen_khoan"], d["the"],
            d["tien_mat"], d["chi_tien_mat"]["tong_chi"], d["tien_mat_phai_con"], d["tien_mat_dem"],
            d["chenh_lech_quy"], d["so_khach"], d["tb_khach"], d["pos_tong_mon"]] for d in dt["ngay"]])

    ds_kho = [d for d in kq["kho"].dong if kq["tu"] <= d.ngay <= den]
    sheet("NXT tổng hợp", f"NHẬP – XUẤT – TỒN – HAO HỤT – {ten_ky}",
          ["NVL", "Mã", "Nhóm", "ĐVT", "Tồn đầu", "Nhập", "Xuất bán", "Huỷ", "Thất thoát", "Thừa", "Tồn cuối",
           "Giá BQ", "GT tồn đầu", "GT nhập", "GT xuất bán", "GT huỷ", "GT thất thoát", "GT tồn cuối",
           "Giá vốn thực tế", "Hao hụt %", "Số lần kiểm"],
          [[v["ten_nvl"], v["ma_nvl"], v["nhom"], v["don_vi"], v["ton_dau"], v["nhap"], v["xuat_ban"], v["huy"],
            v["that_thoat"], v["thua"], v["ton_cuoi"], v["gia_bq"], v["gia_tri_ton_dau"], v["gia_tri_nhap"],
            v["gia_tri_xuat_ban"], v["gia_tri_huy"], v["gia_tri_that_thoat"], v["gia_tri_ton_cuoi"],
            v["gia_von_thuc_te"], v["ty_le_hao_hut"], v["so_lan_kiem"]]
           for v in sorted(tong_hop_theo_nvl(ds_kho).values(), key=lambda x: (x["nhom"], x["ten_nvl"]))],
          {**{i: so_le for i in range(5, 12)}, 20: pct})
    sheet("NXT chi tiết ngày", f"THẺ KHO THEO NGÀY – {ten_ky}",
          ["Ngày", "NVL", "ĐVT", "Tồn đầu", "Nhập", "Xuất bán", "Huỷ", "Tồn lý thuyết", "Kiểm kê thực tế",
           "Chênh lệch", "Tồn cuối", "Giá BQ", "Giá vốn thực tế"],
          [[f"{d.ngay:%d/%m/%Y}", d.ten_nvl, d.don_vi, d.ton_dau, d.nhap, d.xuat_ban, d.huy, d.ton_ly_thuyet,
            d.thuc_te, d.chenh_lech, d.ton_cuoi, d.gia_bq, d.gia_von_thuc_te]
           for d in sorted(ds_kho, key=lambda x: (x.ngay, x.ten_nvl))],
          {i: so_le for i in range(4, 12)})

    dl = kq["dl"]
    sheet("Mua hàng", f"HOÁ ĐƠN MUA HÀNG – {ten_ky}",
          ["Ngày", "NCC", "Chứng từ", "Mặt hàng", "Mã NVL", "SL", "ĐVT", "Đơn giá", "Thành tiền", "Thanh toán",
           "Hạn TT"],
          [[f"{r['ngay']:%d/%m/%Y}", r["ncc"], r.get("so_chung_tu", ""), r["ten_hang"], r.get("ma_nvl", ""),
            r["so_luong"], r.get("don_vi", ""), r["don_gia"], r["thanh_tien"],
            HINH_THUC[chuan_hinh_thuc(r.get("hinh_thuc_tt", ""))],
            f"{r['han_thanh_toan']:%d/%m/%Y}" if r.get("han_thanh_toan") else ""]
           for r in sorted(trong_khoang(dl["nhap_hang.csv"], kq["tu"], den), key=lambda x: x["ngay"])],
          {6: so_le})
    cn = kq["cong_no"]
    sheet("Công nợ NCC", f"CÔNG NỢ NHÀ CUNG CẤP TẠI {den:%d/%m/%Y}",
          ["NCC", "Chứng từ", "Ngày", "Hạn", "Số tiền", "Đã trả", "Còn nợ", "Quá hạn (ngày)", "Ghi chú"],
          [[h.ncc, h.so_chung_tu, f"{h.ngay:%d/%m/%Y}", f"{h.han:%d/%m/%Y}", h.so_tien, h.da_tra, h.con_lai,
            max(0, (den - h.han).days), ("Nợ đầu kỳ. " if h.dau_ky else "") + h.ghi_chu] for h in cn.hoa_don],
          {8: "0"})
    sheet("Cảnh báo", f"CẢNH BÁO & VIỆC CẦN XỬ LÝ – {ten_ky}", ["Mức", "Loại", "Ngày", "Nội dung"],
          [[{"cao": "CAO", "trung_binh": "TRUNG BÌNH", "thap": "THẤP"}.get(c["muc"], c["muc"]), c["loai"],
            f"{c['ngay']:%d/%m/%Y}", c["noi_dung"]] for c in kq["canh_bao"]])
    wb["Cảnh báo"].column_dimensions["D"].width = 120
    tep.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tep)


def lap_bao_cao(dl: DuLieu, tu: date, den: date, loai: str, thu_muc_ra: Path) -> dict:
    kq = tinh_toan(dl, tu, den)
    ten = f"{den:%Y-%m-%d}" if loai == "ngay" else f"{den:%Y-%m}"
    thu_muc = Path(thu_muc_ra) / ("ngay" if loai == "ngay" else "thang")
    thu_muc.mkdir(parents=True, exist_ok=True)
    md = ra_markdown(kq, loai)
    tep_md = thu_muc / f"bao_cao_{ten}.md"
    tep_xlsx = thu_muc / f"bao_cao_{ten}.xlsx"
    tep_md.write_text(md, encoding="utf-8")
    ra_excel(kq, loai, tep_xlsx)
    return {"ket_qua": kq, "markdown": md, "tep_md": tep_md, "tep_xlsx": tep_xlsx}
